from __future__ import annotations

from pathlib import Path

from ebooklib import epub


SAMPLE_RU = (
    "Экономика в практике менеджера — это умение читать спрос, цену и издержки. "
    "Менеджер принимает решения о продукте, канале и марже, опираясь на факты, а не на лозунги."
)


def write_txt(path: Path, text: str = SAMPLE_RU) -> Path:
    path.write_text(text, encoding="utf-8")
    return path


def write_pdf(path: Path, text: str = SAMPLE_RU) -> Path:
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    fontfile = Path("C:/Windows/Fonts/arial.ttf")
    if fontfile.is_file():
        page.insert_font(fontname="F0", fontfile=str(fontfile))
        page.insert_text((72, 72), text, fontsize=12, fontname="F0")
    else:
        page.insert_text((72, 72), text, fontsize=12)
    doc.save(path)
    doc.close()
    return path


def write_docx(path: Path, text: str = SAMPLE_RU) -> Path:
    from docx import Document

    doc = Document()
    doc.add_heading("Глава 1", level=1)
    doc.add_paragraph(text)
    doc.save(path)
    return path


def write_fb2(path: Path, text: str = SAMPLE_RU) -> Path:
    xml = f"""<?xml version="1.0" encoding="utf-8"?>
<FictionBook xmlns="http://www.gribuser.ru/xml/fictionbook/2.0">
  <description>
    <title-info>
      <book-title>Тестовая книга</book-title>
      <author>
        <first-name>Иван</first-name>
        <last-name>Тестов</last-name>
      </author>
    </title-info>
  </description>
  <body>
    <title><p>Глава 1</p></title>
    <p>{text}</p>
  </body>
  <binary id="cover.jpg" content-type="image/jpeg">AAAA</binary>
</FictionBook>
"""
    path.write_text(xml, encoding="utf-8")
    return path


def write_epub(path: Path, text: str = SAMPLE_RU) -> Path:
    book = epub.EpubBook()
    book.set_identifier("rag-fixture")
    book.set_title("Тестовая книга")
    book.set_language("ru")
    chapter = epub.EpubHtml(title="Глава 1", file_name="chap.xhtml", lang="ru")
    chapter.content = f"<html><body><h1>Глава 1</h1><p>{text}</p></body></html>"
    book.add_item(chapter)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav", chapter]
    epub.write_epub(str(path), book)
    return path
